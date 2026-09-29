"""文档解析行为测试：纯文本编码兜底 + PDF 逐页 + DOCX 段落表格"""
from io import BytesIO

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app.services.document_parser import parse_document


def add_text_page(writer: PdfWriter, text: str):
    """向 PdfWriter 增加一页带简单文本流和字体资源的页面"""
    page = writer.add_blank_page(width=400, height=400)

    page[NameObject("/Resources")] = DictionaryObject({
        NameObject("/Font"): DictionaryObject({
            NameObject("/F1"): DictionaryObject({
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }),
        }),
    })

    stream = DecodedStreamObject()
    stream.set_data(f"BT /F1 14 Tf 50 350 Td ({text}) Tj ET".encode("utf-8"))
    page.replace_contents(stream)
    return page


def make_pdf(text: str) -> bytes:
    writer = PdfWriter()
    add_text_page(writer, text)
    buf = BytesIO()
    writer.write(buf)
    return buf.getvalue()


class TestParseTextDocument:
    def test_parses_utf8_text(self):
        units = parse_document("FastAPI 教程。".encode("utf-8"), "guide.md")
        assert units == [("FastAPI 教程。", None)]

    def test_falls_back_to_gbk_when_not_utf8(self):
        data = "中文技术文档内容".encode("gbk")
        units = parse_document(data, "guide.txt")
        assert units[0][0] == "中文技术文档内容"
        assert units[0][1] is None

    def test_txt_and_md_are_treated_as_plain_text(self):
        units_md = parse_document("a".encode(), "x.MD")
        units_txt = parse_document("a".encode(), "x.TXT")
        assert units_md[0][1] is None and units_txt[0][1] is None


class TestParsePdfDocument:
    def test_extracts_text_with_one_based_page_number(self):
        units = parse_document(make_pdf("Hello FastAPI"), "paper.pdf")

        assert len(units) == 1
        text, page = units[0]
        assert "Hello FastAPI" in text
        assert page == 1

    def test_blank_page_is_skipped(self):
        writer = PdfWriter()
        writer.add_blank_page(width=400, height=400)
        buf = BytesIO()
        writer.write(buf)

        units = parse_document(buf.getvalue(), "blank.pdf")
        assert units == []

    def test_multiple_pages_keep_their_page_numbers(self):
        writer = PdfWriter()
        add_text_page(writer, "Page One")
        add_text_page(writer, "Page Two")
        buf = BytesIO()
        writer.write(buf)

        units = parse_document(buf.getvalue(), "two.pdf")

        assert [page for _, page in units] == [1, 2]
        assert "Page One" in units[0][0] and "Page Two" in units[1][0]


def make_docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    """在内存中构造 .docx：若干段落 + 可选表格（首行为表头）"""
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    if table:
        rows = doc.add_table(rows=len(table), cols=len(table[0]))
        for i, row in enumerate(table):
            for j, value in enumerate(row):
                rows.rows[i].cells[j].text = value
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


class TestParseDocxDocument:
    def test_extracts_paragraphs_as_single_unit_without_page(self):
        data = make_docx(["FastAPI 入门", "异步路由性能更好。"])

        units = parse_document(data, "guide.docx")

        assert len(units) == 1
        text, page = units[0]
        assert "FastAPI 入门" in text and "异步路由性能更好。" in text
        assert page is None

    def test_extracts_table_cells(self):
        data = make_docx(["标题"], table=[["方法", "路径"], ["GET", "/users"]])

        text, _ = parse_document(data, "api.docx")[0]

        assert "方法" in text and "GET" in text and "/users" in text

    def test_uppercase_extension_is_accepted(self):
        data = make_docx(["内容"])
        units = parse_document(data, "Guide.DOCX")
        assert units[0][1] is None

    def test_empty_docx_returns_no_units(self):
        data = make_docx([])
        assert parse_document(data, "blank.docx") == []
